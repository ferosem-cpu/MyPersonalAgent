"""Text extraction for dropped work-order files (PDF, DOCX, plain text).

Used by LocalTools.zan_ingest_work_order_file - deliberately just extraction, no
LLM calls here. The calling LLM does the actual field-parsing from the returned text.

Not handled: scanned/image-only PDFs (no selectable text layer) and standalone image
files (.jpg/.png) - both would need OCR, which isn't wired up. extract_text raises a
clear ZanAppError-style message for those rather than silently returning empty text,
so the agent can tell the user rather than hallucinating fields from nothing.
"""
from __future__ import annotations

from pathlib import Path


class ExtractionError(RuntimeError):
    pass


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in (".txt", ".md", ".csv"):
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix in (".jpg", ".jpeg", ".png", ".webp", ".gif"):
        raise ExtractionError(
            f"'{path.name}' is an image file - OCR isn't wired up yet, so this can't "
            "be read automatically. Ask the user to describe its contents, or forward "
            "a text-based version (PDF/DOCX/TXT) instead."
        )
    raise ExtractionError(f"Unsupported file type '{suffix}' for '{path.name}'.")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages_text = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(pages_text).strip()
    if not text:
        raise ExtractionError(
            f"'{path.name}' has no selectable text (likely a scanned image PDF) - "
            "OCR isn't wired up yet, so this can't be read automatically. Ask the "
            "user to describe its contents or forward a text-based version instead."
        )
    return text


def _extract_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise ExtractionError(f"'{path.name}' appears to be empty.")
    return text
